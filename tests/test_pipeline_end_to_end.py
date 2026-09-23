"""End-to-end tests on INVENTED workbooks (see helpers.py) - no real MPLADS data involved."""

import datetime as dt

import pytest
from helpers import (
    FILE_NAMES, HEADERS, ORPHAN, TITLES, W1, W2, W3, W4, W5, W6, build_raw_dir, read_csv, write_workbook,
)

from data_pipeline.build_dataset import PipelineConfig, main, run_pipeline
from data_pipeline.excel_io import SchemaError, file_fingerprint


def by_id(rows):
    return {row["work_id"]: row for row in rows}


def test_all_outputs_are_written_and_raw_files_are_untouched(pipeline):
    expected = [
        "works_master.csv", "expenditure_transactions.csv", "expenditure_by_work.csv", "mp_allocated_limits.csv",
        "calamity_consents.csv", "malformed_ids.csv", "work_id_duplicates.csv", "work_conflicts.csv",
        "join_unmatched_ids.csv", "pipeline_summary.json", "build_dataset.log",
    ]
    for name in expected:
        assert (pipeline.processed / name).is_file(), name
    for name in ("DATA_QUALITY_REPORT.md", "DATA_DICTIONARY.md"):
        assert (pipeline.docs / name).is_file(), name
    assert pipeline.stats.run.raw_unchanged is True
    for key, profile in pipeline.stats.profiles.items():
        assert file_fingerprint(pipeline.raw / FILE_NAMES[key])["sha256"] == profile.fingerprint["sha256"]
    assert not (pipeline.out / "data" / "interim" / "staging.sqlite").exists()  # scratch DB is cleaned up


def test_master_has_one_row_per_valid_work_id_from_recommended_sanctioned_completed(pipeline):
    rows = read_csv(pipeline.processed / "works_master.csv")
    ids = [r["work_id"] for r in rows]
    assert sorted(ids) == sorted([W1, W2, W3, W4, W5, W6])
    assert len(set(ids)) == len(ids)
    assert ORPHAN not in ids  # expenditure-only ID has no master row


def test_master_values_are_normalised_from_mixed_cell_types(pipeline):
    w1 = by_id(read_csv(pipeline.processed / "works_master.csv"))[W1]
    assert w1["recommended_amount"] == "497185"          # numeric cell
    assert w1["sanction_date"] == "2024-07-09"           # real datetime cell in Recommended, text in Sanctioned
    assert w1["recommended_date"] == "2024-07-08"        # DD-Mon-YYYY text
    assert w1["completion_date"] == "2024-09-05" and w1["completed_amount_disbursed"] == "448127"
    assert w1["work"] == "Construction of buildings for community cultural activities"
    assert (w1["id_mp_code"], w1["id_financial_year"]) == ("MP620", "2024-2025")
    assert (w1["in_recommended"], w1["in_sanctioned"], w1["in_completed"]) == ("1", "1", "1")


def test_source_rows_point_back_to_the_raw_workbooks(pipeline):
    rows = by_id(read_csv(pipeline.processed / "works_master.csv"))
    assert (rows[W1]["recommended_source_row"], rows[W1]["sanctioned_source_row"], rows[W1]["completed_source_row"]) == ("3", "3", "3")
    assert rows[W6]["completed_source_row"] == "5" and rows[W6]["recommended_source_row"] == ""
    assert rows[W3]["recommended_source_row"] == "6"      # blank row 5 kept the numbering aligned


def test_sanctioned_values_are_preferred_and_conflicts_are_reported_not_hidden(pipeline):
    master = by_id(read_csv(pipeline.processed / "works_master.csv"))
    assert master[W2]["mp_name"] == "Someone Else"                 # sanctioned wins
    assert master[W2]["conflict_fields"] == "mp_name"
    assert master[W2]["state"] == "KARNATAKA"                      # format-only difference: not a conflict
    assert master[W3]["sanction_date"] == "2024-09-24"             # sanctioned date preferred over 23-Sep in Recommended
    assert master[W3]["conflict_fields"] == "sanction_date"
    conflicts = read_csv(pipeline.processed / "work_conflicts.csv")
    assert {(c["work_id"], c["field"]) for c in conflicts} == {(W2, "mp_name"), (W3, "sanction_date")}
    mp = next(c for c in conflicts if c["field"] == "mp_name")
    assert (mp["chosen_source"], mp["recommended_value"], mp["sanctioned_value"]) == ("sanctioned", "Pralhad Venkatesh Joshi", "Someone Else")
    stats = pipeline.stats.conflicts
    assert stats.works_with_conflicts == 2 and stats.format_only["state"] == 1


def test_repeated_work_ids_keep_the_first_row_and_every_repeat_is_logged(pipeline):
    w3 = by_id(read_csv(pipeline.processed / "works_master.csv"))[W3]
    assert w3["recommended_row_count"] == "3"
    assert w3["recommended_amount"] == "450000"                    # first row, not the 460000 repeat
    dups = read_csv(pipeline.processed / "work_id_duplicates.csv")
    assert [(d["source_row"], d["kept_source_row"], d["relation"], d["differing_fields"]) for d in dups] == [
        ("7", "6", "conflicting", "recommended_amount"),
        ("8", "6", "identical", ""),
    ]
    stats = pipeline.stats
    assert stats.duplicates.extra_rows["recommended"] == 2
    assert (stats.duplicates.identical["recommended"], stats.duplicates.conflicting["recommended"]) == (1, 1)
    assert stats.profiles["recommended"].exact_duplicate_rows == 1
    assert stats.profiles["expenditure"].exact_duplicate_rows == 1


def test_malformed_ids_are_logged_with_reasons_and_never_silently_dropped(pipeline):
    log = read_csv(pipeline.processed / "malformed_ids.csv")
    assert [(r["source"], r["source_row"], r["reason"]) for r in log] == [
        ("recommended", "10", "wrong_segment_count"),
        ("sanctioned", "7", "bad_financial_year"),
        ("expenditure", "9", "wrong_segment_count"),
    ]
    assert log[1]["raw_value"].startswith("WS/ MP620/2024-25/133001")   # original text preserved
    for key, count in (("recommended", 1), ("sanctioned", 1), ("completed", 0), ("expenditure", 1)):
        assert pipeline.stats.profiles[key].id.malformed == count
    master_ids = {r["work_id"] for r in read_csv(pipeline.processed / "works_master.csv")}
    assert "WS/MP620/2024-25/133001" not in master_ids
    # the payment row with the bad ID stays in the transactions file, with an empty work_id
    transactions = read_csv(pipeline.processed / "expenditure_transactions.csv")
    assert len(transactions) == 7 and transactions[-1]["work_id"] == ""


def test_expenditure_transactions_are_one_row_per_payment_and_flag_repeats(pipeline):
    rows = read_csv(pipeline.processed / "expenditure_transactions.csv")
    assert [r["source_row"] for r in rows] == ["3", "4", "5", "6", "7", "8", "9"]
    assert [r["duplicate_record"] for r in rows] == ["0", "0", "0", "1", "0", "0", "0"]   # Sr. No. never matters
    assert rows[1]["fund_disbursed_amount"] == "150000"        # "1,50,000" (Indian grouping)
    assert rows[2]["fund_disbursed_amount"] == ""              # "abc" is invalid -> empty, counted in the report
    assert rows[5]["fund_disbursed_amount"] == "-500"          # negatives are kept and counted
    assert rows[0]["expenditure_date"] == "2024-08-01" and rows[0]["work_id"] == W1
    profile = pipeline.stats.profiles["expenditure"].columns["fund_disbursed_amount"]
    assert (profile.invalid, profile.negatives) == (1, 1)


def test_expenditure_by_work_aggregates(pipeline):
    agg = by_id(read_csv(pipeline.processed / "expenditure_by_work.csv"))
    assert list(agg) == [W1, W3, ORPHAN]
    assert agg[W1] == {
        "work_id": W1, "total_disbursed": "350000", "payment_count": "4",
        "first_expenditure_date": "2024-08-01", "last_expenditure_date": "2024-09-01",
        "vendor_count": "2",                      # "ABC Infra" and "abc  infra" fold to one vendor
        "amount_missing_count": "1", "date_missing_count": "0", "negative_amount_count": "0",
        "duplicate_record_count": "1",
    }
    assert agg[W3]["total_disbursed"] == "400000" and agg[W3]["payment_count"] == "1"
    assert agg[ORPHAN]["total_disbursed"] == "-500" and agg[ORPHAN]["negative_amount_count"] == "1"


def test_join_analysis_counts(pipeline):
    j = pipeline.stats.join
    assert dict(j.unique) == {"R": 4, "S": 4, "C": 3, "E": 3}
    assert (j.union_ids, j.master_rows) == (7, 6)
    assert j.patterns[(True, True, True, True)] == 2
    for pattern in ((True, True, False, False), (True, False, False, False), (False, True, False, False),
                    (False, False, True, False), (False, False, False, True)):
        assert j.patterns[pattern] == 1
    assert (j.pair_total[("S", "R")], j.pair_match[("S", "R")]) == (4, 3)
    assert (j.pair_total[("C", "S")], j.pair_match[("C", "S")]) == (3, 2)
    assert (j.pair_total[("E", "M")], j.pair_match[("E", "M")]) == (3, 2)
    unmatched = {(r["issue"], r["work_id"]) for r in read_csv(pipeline.processed / "join_unmatched_ids.csv")}
    assert unmatched == {
        ("sanctioned_not_in_recommended", W5),
        ("completed_not_in_sanctioned", W6),
        ("completed_without_expenditure", W6),
        ("expenditure_not_in_sanctioned", ORPHAN),
        ("expenditure_not_in_master", ORPHAN),
        ("recommended_with_sanction_date_not_in_sanctioned", W4),
    }


def test_impossible_date_and_amount_relationships_are_counted(pipeline):
    checks = pipeline.stats.checks

    def outcome(name):
        return checks.evaluated[name], checks.hits[name]

    assert outcome("sanction_before_recommendation") == (5, 0)
    assert outcome("completion_before_sanction") == (2, 1)
    assert outcome("completion_before_recommendation") == (2, 1)
    assert outcome("expenditure_before_recommendation") == (2, 0)
    assert outcome("expenditure_before_sanction") == (2, 1)
    assert outcome("completed_amount_exceeds_sanction") == (2, 1)
    assert outcome("expenditure_total_exceeds_sanction") == (2, 0)
    assert outcome("sanction_amount_differs_from_recommended") == (3, 1)
    assert outcome("id_fy_differs_from_sanction_fy") == (5, 1)
    assert checks.samples["completion_before_sanction"] == [W3]


def test_reports_contain_real_figures_and_no_placeholders(pipeline):
    quality = (pipeline.docs / "DATA_QUALITY_REPORT.md").read_text(encoding="utf-8")
    dictionary = (pipeline.docs / "DATA_DICTIONARY.md").read_text(encoding="utf-8")
    for text in (quality, dictionary):
        assert "TBD" not in text and "PARTIAL RUN" not in text
    for heading in ("## 4. Work ID extraction", "## 5. Duplicates", "## 6. Join analysis", "## 7. Conflicting values",
                    "## 8. Column-level quality", "## 10. Date / amount relationships"):
        assert heading in quality
    assert "wrong_segment_count" in quality and "sanctioned_not_in_recommended" in quality
    assert "Raw workbooks unchanged (SHA-256 before == after) | yes" in quality
    assert "| `work_id` |" in dictionary and W1 in dictionary          # example value measured from the output
    assert "**6 rows x 27 columns.**" in dictionary
    assert "Header found in workbook" in dictionary and "`RECOMMENDED AMOUNT   ( \u20b9 )`" in dictionary


def test_demo_subset_is_a_consistent_slice_of_the_processed_data(pipeline):
    master = read_csv(pipeline.demo / "works_master_demo.csv")
    by_work = read_csv(pipeline.demo / "expenditure_by_work_demo.csv")
    transactions = read_csv(pipeline.demo / "expenditure_transactions_demo.csv")
    assert sorted(r["work_id"] for r in master) == sorted(r["work_id"] for r in by_work) == [W1, W3]
    assert len(transactions) == 5 and {t["work_id"] for t in transactions} == {W1, W3}
    full = by_id(read_csv(pipeline.processed / "works_master.csv"))
    assert all(row == full[row["work_id"]] for row in master)     # real rows, not synthetic ones
    assert (pipeline.demo / "README.md").is_file()


def test_output_is_reproducible_byte_for_byte(tmp_path, raw_dir):
    def run(name):
        out = tmp_path / name
        run_pipeline(PipelineConfig(raw_dir=raw_dir, output_root=out, as_of=dt.date(2026, 9, 20), demo_size=2,
                                    demo_seed="s", progress_every=1000))
        return out

    first, second = run("a"), run("b")
    csv_names = sorted(p.name for p in (first / "data" / "processed").glob("*.csv"))
    assert csv_names and csv_names == sorted(p.name for p in (second / "data" / "processed").glob("*.csv"))
    for name in csv_names:
        assert (first / "data" / "processed" / name).read_bytes() == (second / "data" / "processed" / name).read_bytes(), name
    for name in ("works_master_demo.csv", "expenditure_transactions_demo.csv", "expenditure_by_work_demo.csv"):
        assert (first / "data" / "demo" / name).read_bytes() == (second / "data" / "demo" / name).read_bytes(), name


def test_limit_rows_gives_a_partial_run_that_says_so(tmp_path, raw_dir):
    out = tmp_path / "out"
    stats = run_pipeline(PipelineConfig(raw_dir=raw_dir, output_root=out, as_of=dt.date(2026, 9, 20), limit_rows=3,
                                        demo_size=0, progress_every=1000))
    assert stats.profiles["recommended"].data_rows == 3 and stats.profiles["recommended"].limit_hit
    assert "PARTIAL RUN" in (out / "docs" / "DATA_QUALITY_REPORT.md").read_text(encoding="utf-8")
    assert list((out / "data" / "demo").glob("*.csv")) == []           # demo-size 0 writes no demo files


def test_extra_sheets_show_up_in_the_headline_findings(tmp_path, raw_dir):
    from openpyxl import load_workbook

    path = raw_dir / FILE_NAMES["completed"]
    workbook = load_workbook(path)
    workbook.create_sheet("Sheet2").append(["continuation?"])
    workbook.save(path)
    out = tmp_path / "out"
    run_pipeline(PipelineConfig(raw_dir=raw_dir, output_root=out, as_of=dt.date(2026, 9, 20), demo_size=0))
    report = (out / "docs" / "DATA_QUALITY_REPORT.md").read_text(encoding="utf-8")
    assert "**Unread sheets** - Completed" in report and "Sheet2" in report


def test_a_missing_workbook_stops_with_a_clear_error(tmp_path, raw_dir):
    (raw_dir / FILE_NAMES["completed"]).unlink()
    with pytest.raises(FileNotFoundError, match="Works Completed.xlsx"):
        run_pipeline(PipelineConfig(raw_dir=raw_dir, output_root=tmp_path / "out", as_of=dt.date(2026, 9, 20)))


def test_a_workbook_with_a_missing_column_fails_loudly_instead_of_guessing(tmp_path):
    raw = build_raw_dir(tmp_path / "raw")
    headers = [h for h in HEADERS["sanctioned"] if h != "Work Status"]
    write_workbook(raw / FILE_NAMES["sanctioned"], TITLES["sanctioned"], headers, [])
    with pytest.raises(SchemaError, match="work_status"):
        run_pipeline(PipelineConfig(raw_dir=raw, output_root=tmp_path / "out", as_of=dt.date(2026, 9, 20)))


def test_cli_exit_codes(tmp_path, raw_dir):
    out = tmp_path / "out"
    args = ["--raw-dir", str(raw_dir), "--output-root", str(out), "--as-of", "2026-09-20", "--demo-size", "0"]
    assert main(args) == 0
    (raw_dir / FILE_NAMES["calamity"]).unlink()
    assert main(args) == 2
    build_raw_dir(raw_dir)
    write_workbook(raw_dir / FILE_NAMES["allocation"], TITLES["allocation"], ["Nothing", "Useful"], [])
    assert main(args) == 3
