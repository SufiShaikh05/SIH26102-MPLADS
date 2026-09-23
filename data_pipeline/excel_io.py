"""Streaming, read-only access to the raw workbooks.

Nothing here ever writes to a workbook.  ``openpyxl`` read-only mode parses the
sheet XML incrementally, so only one row is materialised at a time (the shared
string table is the main fixed cost).
"""

from __future__ import annotations

import hashlib
import itertools
import warnings
from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

try:  # private module: only used for an optional speed-up, never required
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
except ImportError:  # pragma: no cover - other openpyxl layouts simply skip the optimisation
    ReadOnlyWorksheet = None  # type: ignore[assignment,misc]

from data_pipeline.schemas import SourceSpec, normalize_header

HEADER_SCAN_ROWS = 25


class SchemaError(RuntimeError):
    """The workbook does not contain the columns the pipeline expects."""


@dataclass(slots=True)
class SheetLayout:
    sheet_name: str
    header_row: int                 # 1-based Excel row of the header
    headers: list[str]              # header cells exactly as found
    positions: dict[str, int]       # canonical column -> 0-based position
    header_for: dict[str, str]      # canonical column -> header text as found
    unmapped_headers: list[str]     # headers present in the file but not declared in the spec
    duplicate_headers: list[str]    # repeated headers (first one wins)
    title: str | None               # e.g. "Works Recommended" (row above the header)
    other_sheets: list[str] = field(default_factory=list)   # further sheets in the file - NOT read


@contextmanager
def _without_dimension_scan() -> Iterator[None]:
    """Skip openpyxl's eager ``<dimension>`` lookup while a workbook is opened.

    If a sheet has no ``<dimension>`` element (the spike found exactly that in the exports),
    openpyxl parses the *entire* sheet XML just to discover it - a full extra pass over an
    85 MB workbook.  We throw the size away anyway (see ``reset_dimensions`` below), so the
    scan is skipped.  Best effort: if the hook is missing the reader just runs slower.
    """
    original = getattr(ReadOnlyWorksheet, "_get_size", None) if ReadOnlyWorksheet is not None else None
    if original is None:
        yield
        return
    ReadOnlyWorksheet._get_size = lambda self: None  # type: ignore[method-assign]
    try:
        yield
    finally:
        ReadOnlyWorksheet._get_size = original  # type: ignore[method-assign]


def file_fingerprint(path: Path) -> dict[str, object]:
    """Size, modification time and SHA-256 - identifies the exact snapshot processed."""
    stat = path.stat()
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return {
        "size_bytes": stat.st_size,
        "modified_utc": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(timespec="seconds"),
        "sha256": digest.hexdigest(),
    }


def _alias_map(spec: SourceSpec) -> dict[str, str]:
    return {normalize_header(alias): col.name for col in spec.columns for alias in col.aliases}


def detect_layout(spec: SourceSpec, rows: Sequence[Sequence[object]], sheet_name: str) -> SheetLayout | None:
    """Find the header row among the first rows (title rows above it are skipped)."""
    aliases = _alias_map(spec)
    required = {c.name for c in spec.columns if c.required}
    for index, row in enumerate(rows):
        positions: dict[str, int] = {}
        header_for: dict[str, str] = {}
        unmapped: list[str] = []
        duplicates: list[str] = []
        for pos, cell in enumerate(row):
            if cell is None or not str(cell).strip():
                continue
            canonical = aliases.get(normalize_header(cell))
            if canonical is None:
                unmapped.append(str(cell).strip())
            elif canonical in positions:
                duplicates.append(str(cell).strip())
            else:
                positions[canonical] = pos
                header_for[canonical] = str(cell).strip()
        if required <= positions.keys():
            title = next(
                (str(c).strip() for earlier in rows[:index] for c in earlier if c is not None and str(c).strip()),
                None,
            )
            headers = [str(c).strip() if c is not None else "" for c in row]
            return SheetLayout(sheet_name, index + 1, headers, positions, header_for, unmapped, duplicates, title)
    return None


def _diagnose(spec: SourceSpec, rows: Sequence[Sequence[object]]) -> str:
    aliases = _alias_map(spec)
    required = {c.name for c in spec.columns if c.required}
    best_missing, best_index = sorted(required), None
    for index, row in enumerate(rows):
        found = {aliases.get(normalize_header(c)) for c in row if c is not None and str(c).strip()}
        missing = sorted(required - found)
        if len(missing) < len(best_missing):
            best_missing, best_index = missing, index
    preview = [[str(c)[:30] for c in row if c is not None][:12] for row in rows[:4]]
    where = f"closest header row: {best_index + 1}" if best_index is not None else "no header-like row found"
    return f"{where}; missing required columns: {best_missing}; first rows: {preview}"


class SourceReader:
    """Context manager yielding ``(excel_row_number, {canonical_column: raw_value})``."""

    def __init__(self, spec: SourceSpec, path: Path) -> None:
        self.spec = spec
        self.path = Path(path)
        self.layout: SheetLayout | None = None
        self._wb = None
        self._raw_iter: Iterator[Sequence[object]] | None = None
        self._numbered: Iterator[tuple[int, Sequence[object]]] | None = None

    def __enter__(self) -> "SourceReader":
        with warnings.catch_warnings(), _without_dimension_scan():
            warnings.simplefilter("ignore")  # openpyxl warns about missing default styles etc.
            self._wb = load_workbook(self.path, read_only=True, data_only=True)
        problems: list[str] = []
        for ws in self._wb.worksheets:
            # Some exporters write a wrong/missing <dimension>; ignore it and stream what is there.
            if hasattr(ws, "reset_dimensions"):
                ws.reset_dimensions()
            rows = iter(ws.iter_rows(values_only=True))
            head = list(itertools.islice(rows, HEADER_SCAN_ROWS))
            layout = detect_layout(self.spec, head, ws.title)
            if layout is not None:
                layout.other_sheets = [name for name in self._wb.sheetnames if name != ws.title]
                self.layout = layout
                self._raw_iter = rows
                tail = itertools.chain(head[layout.header_row:], rows)
                self._numbered = enumerate(tail, start=layout.header_row + 1)
                return self
            problems.append(f"sheet {ws.title!r}: {_diagnose(self.spec, head)}")
            getattr(rows, "close", lambda: None)()
        self._wb.close()
        self._wb = None
        raise SchemaError(f"{self.path.name}: expected columns not found. " + " | ".join(problems))

    def __exit__(self, *exc_info: object) -> None:
        close = getattr(self._raw_iter, "close", None)
        if close:
            close()
        if self._wb is not None:
            self._wb.close()  # releases the file handle (matters on Windows)
            self._wb = None

    def rows(self) -> Iterator[tuple[int, dict[str, object]]]:
        assert self.layout is not None and self._numbered is not None, "use SourceReader as a context manager"
        pairs = tuple(self.layout.positions.items())
        for row_no, values in self._numbered:
            width = len(values)  # read-only rows can be shorter than the header
            yield row_no, {name: (values[pos] if pos < width else None) for name, pos in pairs}
